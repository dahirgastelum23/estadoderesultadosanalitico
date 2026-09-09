from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
import os


# ==========================================
# FUNCIÓN PARA PEDIR NÚMEROS
# ==========================================

def pedir_numero(mensaje):
    while True:
        try:
            entrada = input(mensaje).strip()
            entrada = entrada.replace(",", "")
            valor = float(entrada)

            if valor < 0:
                print("El valor no puede ser negativo.")
                continue

            return valor

        except ValueError:
            print("Ingresa un número válido.")
            print("Ejemplo: 500,000 o 1,250,500.50")


# ==========================================
# ENCABEZADO
# ==========================================

print("\n==========================================")
print("      ESTADO DE RESULTADOS ANALÍTICO")
print("==========================================\n")


# ==========================================
# DATOS GENERALES
# ==========================================

empresa = input("Nombre de la empresa: ")
periodo = input("Periodo: ")


# ==========================================
# VENTAS
# ==========================================

print("\n--- VENTAS ---")

ventas = pedir_numero("Ventas: ")
devoluciones_ventas = pedir_numero("Devoluciones sobre ventas: ")
rebajas_ventas = pedir_numero("Rebajas sobre ventas: ")
descuentos_ventas = pedir_numero("Descuentos sobre ventas: ")
bonificaciones_ventas = pedir_numero("Bonificaciones sobre ventas: ")


# ==========================================
# COMPRAS
# ==========================================

print("\n--- COMPRAS ---")

compras = pedir_numero("Compras: ")
gastos_compras = pedir_numero("Gastos sobre compras: ")
devoluciones_compras = pedir_numero("Devoluciones sobre compras: ")
rebajas_compras = pedir_numero("Rebajas sobre compras: ")
descuentos_compras = pedir_numero("Descuentos sobre compras: ")
bonificaciones_compras = pedir_numero("Bonificaciones sobre compras: ")


# ==========================================
# INVENTARIOS
# ==========================================

print("\n--- INVENTARIOS ---")

inventario_inicial = pedir_numero("Inventario inicial: ")
inventario_final = pedir_numero("Inventario final: ")


# ==========================================
# GASTOS DE OPERACIÓN
# ==========================================

print("\n--- GASTOS DE OPERACIÓN ---")

gastos_venta = pedir_numero("Gastos de venta: ")
gastos_administracion = pedir_numero("Gastos de administración: ")
depreciaciones = pedir_numero("Depreciaciones: ")


# ==========================================
# OTROS PRODUCTOS Y GASTOS
# ==========================================

print("\n--- OTROS PRODUCTOS Y GASTOS ---")

productos_financieros = pedir_numero("Productos financieros: ")
otros_productos = pedir_numero("Otros productos: ")
gastos_financieros = pedir_numero("Gastos financieros: ")
otros_gastos = pedir_numero("Otros gastos: ")


# ==========================================
# CÁLCULOS
# ==========================================

# ------------------------------------------
# VENTAS NETAS
# ------------------------------------------

ventas_netas = (
    ventas
    - devoluciones_ventas
    - rebajas_ventas
    - descuentos_ventas
    - bonificaciones_ventas
)


# ------------------------------------------
# COMPRAS TOTALES
# ------------------------------------------

compras_totales = (
    compras
    + gastos_compras
)


# ------------------------------------------
# COMPRAS NETAS
# ------------------------------------------

compras_netas = (
    compras_totales
    - devoluciones_compras
    - rebajas_compras
    - descuentos_compras
    - bonificaciones_compras
)


# ------------------------------------------
# MERCANCÍA DISPONIBLE
# ------------------------------------------

mercancia_disponible = (
    inventario_inicial
    + compras_netas
)


# ------------------------------------------
# COSTO DE VENTAS
# ------------------------------------------

costo_ventas = (
    mercancia_disponible
    - inventario_final
)


# ------------------------------------------
# UTILIDAD BRUTA
# ------------------------------------------

utilidad_bruta = (
    ventas_netas
    - costo_ventas
)


# ------------------------------------------
# TOTAL GASTOS DE OPERACIÓN
# ------------------------------------------

total_gastos_operacion = (
    gastos_venta
    + gastos_administracion
    + depreciaciones
)


# ------------------------------------------
# UTILIDAD DE OPERACIÓN
# ------------------------------------------

utilidad_operacion = (
    utilidad_bruta
    - total_gastos_operacion
)


# ------------------------------------------
# RESULTADO ANTES DE IMPUESTOS
# ------------------------------------------

resultado_antes_impuestos = (
    utilidad_operacion
    + productos_financieros
    + otros_productos
    - gastos_financieros
    - otros_gastos
)


# ==========================================
# PTU E ISR
# ==========================================

# Si existe utilidad, se calculan.
# Si existe pérdida, ambos quedan en $0.

if resultado_antes_impuestos > 0:

    ptu = resultado_antes_impuestos * 0.10
    isr = resultado_antes_impuestos * 0.30

else:

    ptu = 0
    isr = 0


# ------------------------------------------
# RESULTADO NETO
# ------------------------------------------

utilidad_neta = (
    resultado_antes_impuestos
    - ptu
    - isr
)


# ==========================================
# DETERMINAR NOMBRE DEL RESULTADO FINAL
# ==========================================

if utilidad_neta > 0:

    nombre_resultado = "UTILIDAD NETA"

elif utilidad_neta < 0:

    nombre_resultado = "PÉRDIDA NETA"

else:

    nombre_resultado = "RESULTADO NETO"


# ==========================================
# MOSTRAR RESULTADOS
# ==========================================

print("\n==========================================")
print("              RESULTADOS")
print("==========================================")

print(f"Ventas netas:              ${ventas_netas:,.2f}")
print(f"Compras totales:           ${compras_totales:,.2f}")
print(f"Compras netas:             ${compras_netas:,.2f}")
print(f"Mercancía disponible:      ${mercancia_disponible:,.2f}")
print(f"Costo de ventas:           ${costo_ventas:,.2f}")
print(f"Utilidad bruta:            ${utilidad_bruta:,.2f}")
print(f"Gastos de operación:       ${total_gastos_operacion:,.2f}")
print(f"Utilidad de operación:     ${utilidad_operacion:,.2f}")
print(f"Resultado antes impuestos: ${resultado_antes_impuestos:,.2f}")
print(f"PTU (10%):                 ${ptu:,.2f}")
print(f"ISR (30%):                 ${isr:,.2f}")
print(f"{nombre_resultado}:             ${utilidad_neta:,.2f}")


# ==========================================
# CREAR LIBRO DE EXCEL
# ==========================================

wb = Workbook()

ws = wb.active
ws.title = "Estado de Resultados"


# ==========================================
# ESTILOS
# ==========================================

titulo = Font(
    bold=True,
    size=16
)

seccion = Font(
    bold=True,
    size=11
)

resultado = Font(
    bold=True
)

formula_font = Font(
    italic=True,
    size=9
)

centrado = Alignment(
    horizontal="center",
    vertical="center"
)

derecha = Alignment(
    horizontal="right"
)

thin = Side(
    style="thin"
)

border = Border(
    top=thin,
    bottom=thin
)


# ==========================================
# TÍTULO
# ==========================================

ws.merge_cells("A1:C1")

ws["A1"] = "ESTADO DE RESULTADOS ANALÍTICO"

ws["A1"].font = titulo
ws["A1"].alignment = centrado


# ==========================================
# INFORMACIÓN GENERAL
# ==========================================

ws["A2"] = "Empresa:"
ws["B2"] = empresa

ws["A3"] = "Periodo:"
ws["B3"] = periodo

ws["A2"].font = resultado
ws["A3"].font = resultado


# ==========================================
# ENCABEZADOS
# ==========================================

ws["A5"] = "Concepto"
ws["B5"] = "Importe"
ws["C5"] = "Fórmula"

for celda in ws[5]:

    celda.font = resultado
    celda.alignment = centrado
    celda.border = border


# ==========================================
# VENTAS
# ==========================================

ws["A6"] = "VENTAS"
ws["A6"].font = seccion

ws["A7"] = "Ventas"
ws["B7"] = ventas

ws["A8"] = "Devoluciones sobre ventas"
ws["B8"] = devoluciones_ventas

ws["A9"] = "Rebajas sobre ventas"
ws["B9"] = rebajas_ventas

ws["A10"] = "Descuentos sobre ventas"
ws["B10"] = descuentos_ventas

ws["A11"] = "Bonificaciones sobre ventas"
ws["B11"] = bonificaciones_ventas

ws["A12"] = "Ventas netas"
ws["B12"] = "=B7-B8-B9-B10-B11"
ws["C12"] = "=B7-B8-B9-B10-B11"


# ==========================================
# COMPRAS
# ==========================================

ws["A14"] = "COMPRAS"
ws["A14"].font = seccion

ws["A15"] = "Compras"
ws["B15"] = compras

ws["A16"] = "Gastos sobre compras"
ws["B16"] = gastos_compras

ws["A17"] = "Compras totales"
ws["B17"] = "=B15+B16"
ws["C17"] = "=B15+B16"

ws["A18"] = "Devoluciones sobre compras"
ws["B18"] = devoluciones_compras

ws["A19"] = "Rebajas sobre compras"
ws["B19"] = rebajas_compras

ws["A20"] = "Descuentos sobre compras"
ws["B20"] = descuentos_compras

ws["A21"] = "Bonificaciones sobre compras"
ws["B21"] = bonificaciones_compras

ws["A22"] = "Compras netas"
ws["B22"] = "=B17-B18-B19-B20-B21"
ws["C22"] = "=B17-B18-B19-B20-B21"


# ==========================================
# COSTO DE VENTAS
# ==========================================

ws["A24"] = "COSTO DE VENTAS"
ws["A24"].font = seccion

ws["A25"] = "Inventario inicial"
ws["B25"] = inventario_inicial

ws["A26"] = "Compras netas"
ws["B26"] = "=B22"

ws["A27"] = "Mercancía disponible"
ws["B27"] = "=B25+B26"
ws["C27"] = "=B25+B26"

ws["A28"] = "Inventario final"
ws["B28"] = inventario_final

ws["A29"] = "Costo de ventas"
ws["B29"] = "=B27-B28"
ws["C29"] = "=B27-B28"


# ==========================================
# UTILIDAD BRUTA
# ==========================================

ws["A31"] = "Utilidad bruta"

ws["B31"] = "=B12-B29"

ws["C31"] = "=B12-B29"

ws["A31"].font = resultado
ws["B31"].font = resultado


# ==========================================
# GASTOS DE OPERACIÓN
# ==========================================

ws["A33"] = "GASTOS DE OPERACIÓN"
ws["A33"].font = seccion

ws["A34"] = "Gastos de venta"
ws["B34"] = gastos_venta

ws["A35"] = "Gastos de administración"
ws["B35"] = gastos_administracion

ws["A36"] = "Depreciaciones"
ws["B36"] = depreciaciones

ws["A37"] = "Total gastos de operación"

ws["B37"] = "=SUM(B34:B36)"

ws["C37"] = "=SUM(B34:B36)"


# ==========================================
# UTILIDAD DE OPERACIÓN
# ==========================================

ws["A39"] = "Utilidad de operación"

ws["B39"] = "=B31-B37"

ws["C39"] = "=B31-B37"

ws["A39"].font = resultado
ws["B39"].font = resultado


# ==========================================
# OTROS PRODUCTOS Y GASTOS
# ==========================================

ws["A41"] = "OTROS PRODUCTOS Y GASTOS"
ws["A41"].font = seccion

ws["A42"] = "Productos financieros"
ws["B42"] = productos_financieros

ws["A43"] = "Otros productos"
ws["B43"] = otros_productos

ws["A44"] = "Gastos financieros"
ws["B44"] = gastos_financieros

ws["A45"] = "Otros gastos"
ws["B45"] = otros_gastos


# ==========================================
# RESULTADO ANTES DE IMPUESTOS
# ==========================================

ws["A47"] = "Resultado antes de impuestos"

ws["B47"] = "=B39+B42+B43-B44-B45"

ws["C47"] = "=B39+B42+B43-B44-B45"

ws["A47"].font = resultado
ws["B47"].font = resultado


# ==========================================
# PTU E ISR
# ==========================================

ws["A49"] = "PTU E ISR"
ws["A49"].font = seccion


# PTU
# Si hay pérdida, queda en $0.

ws["A50"] = "PTU (10%)"

ws["B50"] = "=MAX(B47*10%,0)"

ws["C50"] = "=MAX(B47*10%,0)"


# ISR
# Si hay pérdida, queda en $0.

ws["A51"] = "ISR (30%)"

ws["B51"] = "=MAX(B47*30%,0)"

ws["C51"] = "=MAX(B47*30%,0)"


# ==========================================
# RESULTADO FINAL
# ==========================================

# El texto cambia automáticamente dependiendo
# del resultado calculado en B53.

ws["A53"] = '=IF(B53>0,"UTILIDAD NETA",IF(B53<0,"PÉRDIDA NETA","RESULTADO NETO"))'

ws["B53"] = "=B47-B50-B51"

ws["C53"] = "=B47-B50-B51"


ws["A53"].font = Font(
    bold=True,
    size=12
)

ws["B53"].font = Font(
    bold=True,
    size=12
)


# ==========================================
# FORMATO MONETARIO
# ==========================================

for row in range(7, 54):

    ws[f"B{row}"].number_format = '$#,##0.00;-$#,##0.00'

    ws[f"B{row}"].alignment = derecha


# ==========================================
# FORMATO DE FÓRMULAS
# ==========================================

for row in range(1, 54):

    ws[f"C{row}"].font = formula_font


# ==========================================
# BORDES EN RESULTADOS
# ==========================================

filas_resultado = [
    12,
    17,
    22,
    27,
    29,
    31,
    37,
    39,
    47,
    50,
    51,
    53
]

for fila in filas_resultado:

    for columna in ["A", "B", "C"]:

        ws[f"{columna}{fila}"].border = Border(
            top=thin,
            bottom=thin
        )


# ==========================================
# ANCHO DE COLUMNAS
# ==========================================

ws.column_dimensions["A"].width = 42
ws.column_dimensions["B"].width = 20
ws.column_dimensions["C"].width = 32


# ==========================================
# CONFIGURACIÓN DE LA HOJA
# ==========================================

ws.freeze_panes = "A6"

ws.print_title_rows = "1:5"

ws.page_setup.orientation = "portrait"

ws.page_setup.fitToWidth = 1

ws.page_setup.fitToHeight = 0

ws.sheet_properties.pageSetUpPr.fitToPage = True


# ==========================================
# GENERAR NOMBRE AUTOMÁTICO
# ==========================================

nombre_base = "estado_resultados"

extension = ".xlsx"

contador = 1

while True:

    if contador == 1:

        nombre_archivo = (
            nombre_base
            + extension
        )

    else:

        nombre_archivo = (
            f"{nombre_base}_{contador}{extension}"
        )

    if not os.path.exists(nombre_archivo):

        break

    contador += 1


# ==========================================
# GUARDAR ARCHIVO
# ==========================================

wb.save(nombre_archivo)


# ==========================================
# MENSAJE FINAL
# ==========================================

print("\n==========================================")
print("       Excel creado correctamente.")
print(f"       Archivo: {nombre_archivo}")
print("==========================================")
